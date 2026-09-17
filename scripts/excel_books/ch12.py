"""Ch12 進階案例：LC-MS/MS 的量測不確定度（基質效應、SIL-IS、1/x 加權校正）。

R 的答案來源：Rscript R/case_lcmsms.R。
Matuszewski 三組實驗的原始 R 檔只直接給定三條校正線的「斜率」(45200/38400/34600)，
沒有留下逐點的濃度/波峰面積數據；為了讓學生在 Excel 用 SLOPE() 練習，
本活頁簿另外設計了「濃度 x、面積 y」的示範數據（零截距線性關係，x 任取 0/2/5/10/20），
其斜率被設計為與 R 案例完全相同的 45200/38400/34600——這點在「說明」頁已明確告知學生。
1/x 加權迴歸的 conc/ratio 兩欄則是直接用 set.seed(11) 重新執行 R 檔取得的逐點數值，
與 R 檔生成的教材數據完全一致。
"""
from openpyxl.styles import Font, Alignment

from excel_common import Book, fx, BOX, FONT, HEAD_FILL, INPUT_FILL

BUDGET_HEADERS = ["來源", "數值", "分布", "除數", "標準不確定度 u", "相對 u(x)/x", "u²占比 (%)"]


def budget_header(ws, row):
    for i, text in enumerate(BUDGET_HEADERS):
        c = ws.cell(row=row, column=1 + i, value=text)
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX


def budget_cell(ws, row, col, value, fmt=None, bold=False, fill=None):
    c = ws.cell(row=row, column=col, value=fx(value) if isinstance(value, str) else value)
    c.font = Font(name=FONT, bold=bold)
    c.border = BOX
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    return c


def build():
    book = Book("ch12_lcmsms_uncertainty.xlsx",
                 "Ch12 進階案例：LC-MS/MS 基質效應、SIL-IS 與 1/x 加權校正的不確定度",
                 "ch12.html")
    book.readme([
        "# 這本活頁簿在做什麼",
        "以蝦仁中氯黴素（LC-MS/MS，MRPL=0.3 μg/kg）為例，重現三個 LC-MS/MS 特有的統計工具：",
        "①Matuszewski 三組實驗（ME%基質效應／RE%萃取回收／PE%流程效率，都用 SLOPE 算斜率再相除）",
        "②同位素內標（SIL-IS）面積比校正——同樣的三個指標，比較補償前後的差異",
        "③1/x 加權迴歸——痕量分析必備，用 SUMPRODUCT 手刻加權最小平方公式（Excel 沒有內建加權迴歸）",
        "最後一張工作表把四個不確定度成分（中間精密度／均勻性／校正殘餘／回收率）合成，並做限量符合性判定。",
        "# 怎麼用",
        "1. 黃色是可以改的數據；改了以後所有公式自動重算。",
        "2.「Matuszewski基質效應」工作表的 x、y 數據是本活頁簿設計的示範數據（零截距直線），斜率設計成與 R 案例的 45200/38400/34600 完全一致——這樣你可以用 Excel 的 SLOPE() practice 算斜率比值，效果與 R 一樣。",
        "3.「加權迴歸1x」工作表可以比較 OLS 與 1/x 加權在低濃度的回算誤差差多大。",
        "4.「不確定度預算與符合性」工作表最下面有一張 Monte Carlo 示範表——按 F9 重算會變動，不列入核對。",
        "# 兩個最重要的觀念",
        "① ME%＝(斜率B/斜率A −1)×100 是「電離效率」被基質壓制的程度；RE%＝(斜率C/斜率B −1)×100 才是「萃取回收」——兩者容易搞混，分子分母的順序決定看到的是哪一段流程。",
        "② 1/x 加權不是「加重高濃度點」，而是「讓低濃度點的份量不被高濃度的大誤差蓋過」；痕量分析（誤差與濃度成正比）幾乎都要加權，否則低濃度端會嚴重偏倚。",
        "# 本章沒有對應的資料分析工具箱功能",
        "SLOPE() 本身可以對照工具箱的『迴歸』分析（斜率＝迴歸係數），但 Matuszewski 比值、1/x 加權迴歸、不確定度預算都是逐項代公式，工具箱做不到，本章一律用一般函數現場計算。",
    ])

    # =================================================================
    # 工作表 1：Matuszewski 三組實驗
    # =================================================================
    ws = book.sheet("Matuszewski基質效應", [26, 16, 44, 60])
    book.text(ws, 1, 1, "示範數據：濃度 x 與三組校正線的波峰面積 y（零截距，斜率＝R 案例的值）", bold=True, size=11, color="0F4C81")
    book.header(ws, 2, ["濃度 x (相對單位)"])
    x_rng = book.data(ws, 3, 1, [0, 2, 5, 10, 20], fmt="0")
    book.header(ws, 2, ["yA 溶液標準"], col=2)
    yA_rng = book.data(ws, 3, 2, [0, 90400, 226000, 452000, 904000], fmt="0")
    book.header(ws, 2, ["yB 萃取後添加"], col=3)
    yB_rng = book.data(ws, 3, 3, [0, 76800, 192000, 384000, 768000], fmt="0")
    book.header(ws, 2, ["yC 添加後萃取"], col=4)
    yC_rng = book.data(ws, 3, 4, [0, 69200, 173000, 346000, 692000], fmt="0")

    book.header(ws, 9, ["項目", "結果", "公式", "說明"])
    slope_A = book.calc(ws, 10, "斜率 A（溶液標準）", f"=SLOPE({yA_rng},{x_rng})", "理想電離效率", fmt="0")
    slope_B = book.calc(ws, 11, "斜率 B（萃取後添加）", f"=SLOPE({yB_rng},{x_rng})", "純電離表現（沒有萃取步驟）", fmt="0")
    slope_C = book.calc(ws, 12, "斜率 C（添加後萃取）", f"=SLOPE({yC_rng},{x_rng})", "電離+萃取損失（像真樣品）", fmt="0")
    ME = book.calc(ws, 13, "ME%：基質效應", f"=({slope_B}/{slope_A}-1)*100", "(B/A−1)×100，負值＝離子壓制", fmt="0.0000", key=True)
    RE = book.calc(ws, 14, "RE%：萃取回收", f"=({slope_C}/{slope_B}-1)*100", "(C/B−1)×100", fmt="0.0000", key=True)
    PE = book.calc(ws, 15, "PE%：流程效率", f"=({slope_C}/{slope_A}-1)*100", "(C/A−1)×100，整體流程效率", fmt="0.0000", key=True)

    book.text(ws, 17, 1, "同位素內標（SIL-IS）面積比校正後，重跑三組實驗", bold=True, size=12, color="0F4C81")
    book.header(ws, 18, ["項目", "結果", "公式", "說明"])
    s_sol = book.calc(ws, 19, "面積比斜率 A_IS（溶液標準）", 1.020, "分析物/內標 面積比", fmt="0.0000")
    book.mark_input(ws, s_sol)
    s_post = book.calc(ws, 20, "面積比斜率 B_IS（萃取後添加）", 0.980, "", fmt="0.0000")
    book.mark_input(ws, s_post)
    s_pre = book.calc(ws, 21, "面積比斜率 C_IS（添加後萃取）", 0.972, "", fmt="0.0000")
    book.mark_input(ws, s_pre)
    ME_IS = book.calc(ws, 22, "ME%_IS（補償後）", f"=({s_post}/{s_sol}-1)*100", "從 −15.0% 縮到這個值", fmt="0.0000", key=True)
    RE_IS = book.calc(ws, 23, "RE%_IS（補償後）", f"=({s_pre}/{s_post}-1)*100", "從 −9.9% 縮到這個值", fmt="0.0000", key=True)
    book.text(ws, 25, 1, "殘餘基質效應（本例估 0.08）計入下一張工作表的不確定度預算「校正＋基質殘餘」項。", bold=True)

    # =================================================================
    # 工作表 2：1/x 加權迴歸
    # =================================================================
    ws2 = book.sheet("加權迴歸1x", [30, 16, 16, 16, 16, 46])
    book.text(ws2, 1, 1, "校正數據（與 R 檔 set.seed(11) 完全相同的教材數據）", bold=True, size=12, color="0F4C81")
    book.header(ws2, 2, ["濃度 conc (μg/kg)"])
    conc_rng = book.data(ws2, 3, 1, [0.02, 0.05, 0.10, 0.20, 0.50], fmt="0.00")
    book.header(ws2, 2, ["面積比 ratio"], col=2)
    ratio_rng = book.data(ws2, 3, 2, [0.0393283116, 0.1011189336, 0.1914084760, 0.3888120579, 1.0176663535], fmt="0.0000000")
    book.header(ws2, 2, ["權重 w=1/conc"], col=3)
    for i in range(5):
        row = 3 + i
        c = ws2.cell(row=row, column=3, value=fx(f"=1/A{row}"))
        c.number_format = "0.00"
        c.border = BOX
    w_rng = "C3:C7"

    book.text(ws2, 9, 1, "① 一般最小平方 OLS", bold=True, size=11, color="0F4C81")
    book.header(ws2, 10, ["項目", "結果", "公式", "說明"])
    b1_ols = book.calc(ws2, 11, "b1_OLS：斜率", f"=SLOPE({ratio_rng},{conc_rng})", "", fmt="0.000000", key=True)
    b0_ols = book.calc(ws2, 12, "b0_OLS：截距", f"=INTERCEPT({ratio_rng},{conc_rng})", "", fmt="0.000000", key=True)

    book.text(ws2, 14, 1, "② 1/x 加權迴歸（手刻加權最小平方，全部用 SUMPRODUCT）", bold=True, size=11, color="0F4C81")
    book.header(ws2, 15, ["項目", "結果", "公式", "說明"])
    xbar_w = book.calc(ws2, 16, "x̄_w：加權平均濃度", f"=SUMPRODUCT({w_rng},{conc_rng})/SUM({w_rng})",
                        "Σ(w·x)/Σw", fmt="0.0000000")
    ybar_w = book.calc(ws2, 17, "ȳ_w：加權平均面積比", f"=SUMPRODUCT({w_rng},{ratio_rng})/SUM({w_rng})",
                        "Σ(w·y)/Σw", fmt="0.0000000")
    sxy_w = book.calc(ws2, 18, "S_xy,w：加權交乘和",
                       f"=SUMPRODUCT({w_rng},({conc_rng}-{xbar_w}),({ratio_rng}-{ybar_w}))",
                       "Σw(x−x̄_w)(y−ȳ_w)——SUMPRODUCT 可以直接吃陣列運算，不用 Ctrl+Shift+Enter", fmt="0.0000000")
    sxx_w = book.calc(ws2, 19, "S_xx,w：加權平方和",
                       f"=SUMPRODUCT({w_rng},({conc_rng}-{xbar_w})^2)",
                       "Σw(x−x̄_w)²", fmt="0.0000000")
    b1_w = book.calc(ws2, 20, "b1_w：加權斜率", f"={sxy_w}/{sxx_w}", "S_xy,w / S_xx,w", fmt="0.000000", key=True)
    b0_w = book.calc(ws2, 21, "b0_w：加權截距", f"={ybar_w}-{b1_w}*{xbar_w}", "ȳ_w − b1_w·x̄_w", fmt="0.000000", key=True)

    book.text(ws2, 23, 1, "③ 回算比較表（哪個方法在低濃度偏得少？）", bold=True, size=11, color="0F4C81")
    book.header(ws2, 24, ["名義濃度", "找回濃度_OLS", "OLS回算誤差%", "找回濃度_加權", "加權回算誤差%"])
    for i in range(5):
        row = 25 + i
        crow = 3 + i
        budget_cell(ws2, row, 1, f"=A{crow}", fmt="0.00")
        f_ols = ws2.cell(row=row, column=2, value=fx(f"=(B{crow}-{b0_ols})/{b1_ols}"))
        f_ols.number_format = "0.0000"
        f_ols.border = BOX
        re_ols = ws2.cell(row=row, column=3, value=fx(f"=(B{row}/A{row}-1)*100"))
        re_ols.number_format = "0.0"
        re_ols.border = BOX
        f_w = ws2.cell(row=row, column=4, value=fx(f"=(B{crow}-{b0_w})/{b1_w}"))
        f_w.number_format = "0.0000"
        f_w.border = BOX
        re_w = ws2.cell(row=row, column=5, value=fx(f"=(D{row}/A{row}-1)*100"))
        re_w.number_format = "0.0"
        re_w.border = BOX
    found_w_002 = "D25"

    book.text(ws2, 31, 1, "④ 未知樣品反推（面積比 0.0405）", bold=True, size=11, color="0F4C81")
    book.header(ws2, 32, ["項目", "結果", "公式", "說明"])
    y_obs = book.calc(ws2, 33, "未知樣品面積比 y_obs", 0.0405, "", fmt="0.0000")
    book.mark_input(ws2, y_obs)
    x_ols_pred = book.calc(ws2, 34, "反推濃度（OLS）", f"=({y_obs}-{b0_ols})/{b1_ols}", "", fmt="0.00000")
    x_w_pred = book.calc(ws2, 35, "反推濃度（1/x 加權）", f"=({y_obs}-{b0_w})/{b1_w}", "", fmt="0.00000", key=True)

    book.text(ws2, 37, 1, "Monte Carlo 示範：加權反推濃度的變動範圍（F9 重算會變，不列入核對）", bold=True, size=10, color="64748B")
    book.header(ws2, 38, ["模擬編號", "模擬 y*", "反推 x*（用固定 b1_w/b0_w 近似）"])
    for i in range(30):
        row = 39 + i
        budget_cell(ws2, row, 1, i + 1, fmt="0")
        ysim = ws2.cell(row=row, column=2, value=fx(f"=NORM.INV(RAND(),{y_obs},0.003)"))
        ysim.number_format = "0.0000"
        ysim.border = BOX
        xsim = ws2.cell(row=row, column=3, value=fx(f"=(B{row}-{b0_w})/{b1_w}"))
        xsim.number_format = "0.00000"
        xsim.border = BOX
    book.text(ws2, 70, 1, "完整的加權校正 Monte Carlo（每條模擬校正線都重新配一次）在 R 檔示範，Excel 不易做到逐條重配，這裡只示範「觀測值本身的隨機性」對反推濃度的影響。", size=10, color="64748B")

    # =================================================================
    # 工作表 3：不確定度預算與符合性
    # =================================================================
    ws3 = book.sheet("不確定度預算與符合性", [26, 14, 20, 12, 16, 14, 12, 46])
    book.text(ws3, 1, 1, "四個相對標準不確定度成分（內部驗證法，SANTE 風格）", bold=True, size=12, color="0F4C81")
    budget_header(ws3, 2)
    rows_spec = [
        ("中間精密度", 0.22, "驗證期不同天/人/校正週期的樣品數據"),
        ("校正＋基質殘餘", 0.08, "SIL-IS＋基質匹配校正後的殘餘基質效應"),
        ("回收率偏倚", 0.06, "加標回收平均 90%，平均值的標準誤"),
        ("樣品均勻性", 0.15, "痕量污染物在蝦肉組織的分布（最壞情境模型）"),
    ]
    rel_cells = []
    pct_cells = []
    for i, (name, val, note) in enumerate(rows_spec):
        r = 3 + i
        budget_cell(ws3, r, 1, name)
        budget_cell(ws3, r, 2, 1.0, fmt="0.00")
        budget_cell(ws3, r, 3, note)
        budget_cell(ws3, r, 4, "—")
        u_cell = budget_cell(ws3, r, 5, val, fmt="0.0000", fill=INPUT_FILL)
        rel_cell = budget_cell(ws3, r, 6, f"={u_cell.coordinate}/1", fmt="0.0000")
        pct_cell = budget_cell(ws3, r, 7, None, fmt="0.0")
        rel_cells.append(rel_cell)
        pct_cells.append(pct_cell)

    book.header(ws3, 8, ["項目", "結果", "公式", "說明"])
    rel_uc = book.calc(ws3, 9, "相對合成不確定度 u_c(w)/w",
                        f"=SQRT(SUMSQ({rel_cells[0].coordinate},{rel_cells[1].coordinate},{rel_cells[2].coordinate},{rel_cells[3].coordinate}))",
                        "四項平方和開根號", fmt="0.0000", key=True)
    for pct_cell, rel_cell in zip(pct_cells, rel_cells):
        pct_cell.value = fx(f"={rel_cell.coordinate}^2/{rel_uc}^2*100")

    raw = book.calc(ws3, 11, "儀器讀值（同位素稀釋，已扣除基質） (μg/kg)", 0.19, "", fmt="0.0000")
    book.mark_input(ws3, raw)
    Rec = book.calc(ws3, 12, "Rec：平均回收率", 0.90, "", fmt="0.0000")
    book.mark_input(ws3, Rec)
    result = book.calc(ws3, 13, "w：回收率修正後結果 (μg/kg)", f"={raw}/{Rec}", "raw / Rec", fmt="0.0000", key=True)
    U = book.calc(ws3, 14, "U：擴展不確定度 (k=2, μg/kg)", f"=2*{rel_uc}*{result}", "2 × 相對合成 × w", fmt="0.0000", key=True)
    book.text(ws3, 15, 1, "報告：氯黴素 = (0.21 ± 0.12) μg/kg（k=2）", bold=True)

    book.text(ws3, 17, 1, "符合性判定（保守決策規則）", bold=True, size=12, color="0F4C81")
    book.header(ws3, 18, ["項目", "結果", "公式", "說明"])
    MRL = book.calc(ws3, 19, "MRPL：管制界限 (μg/kg)", 0.3, "", fmt="0.00")
    book.mark_input(ws3, MRL)
    lo = book.calc(ws3, 20, "下限 = w − U", f"={result}-{U}", "", fmt="0.0000")
    hi = book.calc(ws3, 21, "上限 = w + U", f"={result}+{U}", "", fmt="0.0000")
    verdict = book.calc(ws3, 22, "判定",
                         f'=IF({lo}>{MRL},"不符合",IF({hi}<={MRL},"符合","灰色地帶"))',
                         "下限超過限量才判不符合；上限沒超過才判符合；其餘是灰色地帶", fmt="@", key=True)

    book.text(ws3, 24, 1, "Monte Carlo 決策支持示範：真值超過管制界限的機率（F9 重算會變，不列入核對）", bold=True, size=10, color="64748B")
    book.header(ws3, 25, ["模擬編號", "模擬真值 (μg/kg)", "是否超標"])
    for i in range(40):
        row = 26 + i
        budget_cell(ws3, row, 1, i + 1, fmt="0")
        sim = ws3.cell(row=row, column=2, value=fx(f"=NORM.INV(RAND(),{result},{rel_uc}*{result})"))
        sim.number_format = "0.0000"
        sim.border = BOX
        over = ws3.cell(row=row, column=3, value=fx(f'=IF(B{row}>{MRL},1,0)'))
        over.number_format = "0"
        over.border = BOX
    p_over = book.calc(ws3, 67, "估計超標機率 P(真值>MRPL)", "=AVERAGE(C26:C65)",
                        "模擬次數只有 40，數字會隨 F9 大幅跳動；R 用 20 萬次得約 7%——這裡只是示範原理", fmt="0.0%")

    book.check("Matuszewski 斜率 A", "Matuszewski基質效應", slope_A, 45200)
    book.check("ME%（基質效應）", "Matuszewski基質效應", ME, -15.04424779)
    book.check("RE%（萃取回收）", "Matuszewski基質效應", RE, -9.895833333)
    book.check("PE%（流程效率）", "Matuszewski基質效應", PE, -23.45132743)
    book.check("ME%_IS（補償後）", "Matuszewski基質效應", ME_IS, -3.921568627)
    book.check("RE%_IS（補償後）", "Matuszewski基質效應", RE_IS, -0.8163265306)
    book.check("OLS 斜率 b1", "加權迴歸1x", b1_ols, 2.040052862694)
    book.check("OLS 截距 b0", "加權迴歸1x", b0_ols, -0.007302371574)
    book.check("加權斜率 b1_w", "加權迴歸1x", b1_w, 2.008753393)
    book.check("加權截距 b0_w", "加權迴歸1x", b0_w, -0.00185626386)
    book.check("反推濃度（1/x 加權）", "加權迴歸1x", x_w_pred, 0.02108584558)
    book.check("找回濃度_加權 (conc=0.02)", "加權迴歸1x", found_w_002, 0.02050255427)
    book.check("LC-MS/MS 相對合成不確定度", "不確定度預算與符合性", rel_uc, 0.2844292531)
    book.check("LC-MS/MS 結果 w", "不確定度預算與符合性", result, 0.2111111111)
    book.check("LC-MS/MS U (k=2)", "不確定度預算與符合性", U, 0.1200923513)
    return book
