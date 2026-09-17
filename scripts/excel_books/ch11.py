"""Ch11 綜合案例：粗纖維（QUAM Example A6）與標準曲線反推的不確定度（QUAM Appendix E.4）。

R 的答案來源：Rscript R/ch11_capstone.R（案例②粗纖維、案例③標準曲線反推）。
"""
from openpyxl.styles import Font, Alignment

from excel_common import Book, fx, BOX, FONT, HEAD_FILL, INPUT_FILL

BUDGET_HEADERS = ["來源", "數值", "分布", "除數", "標準不確定度 u", "相對 u(x)/x", "u²占比 (%)"]


def budget_header(ws, row):
    for i, text in enumerate(BUDGET_HEADERS):
        c = ws.cell(row=row, column=1 + i, value=text)
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX


def budget_cell(ws, row, col, value, fmt=None, bold=False, fill=None):
    c = ws.cell(row=row, column=col, value=fx(value) if isinstance(value, str) else value)
    c.font = Font(name=FONT, bold=bold)
    c.border = BOX
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    return c


def build():
    book = Book("ch11_fibre_calibration.xlsx",
                 "Ch11 綜合案例：粗纖維（經驗方法）與標準曲線反推的不確定度",
                 "ch11.html")
    book.readme([
        "# 這本活頁簿在做什麼",
        "①飼料粗纖維：一個「結果由方法本身定義」的經驗方法，不確定度直接用協同試驗的再現性標準差 s_R 估（QUAM Example A6）。",
        "②標準曲線反推濃度：把第4章的迴歸與信賴區間串成一條線——用 QUAM 附錄 E.4 的公式，把「殘差、斜率、校正點數、未知樣品重複數」分開拆算。",
        "「粗纖維」工作表：3 個含量層級（2.5%／5%／10%）各自的不確定度預算，看相對不確定度怎麼隨含量遞減。",
        "「標準曲線反推」工作表：SLOPE、INTERCEPT、STEYX、DEVSQ 逐項算出 var(x_pred) 公式的每一項。",
        "# 怎麼用",
        "1. 黃色是可以改的數據；改了以後所有公式自動重算。",
        "2. 試著把「標準曲線反推」工作表裡未知樣品的重複測定次數 p 從 3 改成 1，看 u(x_pred) 變大多少。",
        "3. 把粗纖維的含量層級 s_R 改一改，看相對不確定度 U/含量 怎麼隨含量升高而下降。",
        "4.「對照R答案」工作表核對 Excel 與 R 是否一致。",
        "# 兩個最重要的觀念",
        "① 經驗（實證）方法沒有獨立於方法之外的「真值」，偏倚依定義為零、不修正；但不同實驗室做出來的散布仍要報不確定度，主體是協同試驗的 s_R。",
        "② 標準曲線反推的不確定度在校正範圍「中段最小、兩端最寬」——這就是 (x_pred−x̄)²/S_xx 這一項的幾何意義。",
        "# 本章沒有對應的資料分析工具箱功能",
        "s_R 對含量的關係圖可以用「插入圖表」畫，但預算計算本身是逐項代公式，不是統計檢定；var(x_pred) 公式雖然來自迴歸，仍需拆開手算，工具箱的『迴歸』只給係數，不會算這個反推不確定度公式。",
    ])

    # =================================================================
    # 工作表 1：粗纖維（QUAM Example A6）
    # =================================================================
    ws = book.sheet("粗纖維協同試驗", [26, 14, 14, 14, 14, 46])
    book.text(ws, 1, 1, "協同試驗數據（QUAM Example A6，5 個樣品）", bold=True, size=12, color="0F4C81")
    book.header(ws, 2, ["樣品", "纖維含量 (% m/m)", "s_R 再現性SD", "s_r 重複性SD"])
    names = ["A", "B", "C", "D", "E"]
    for i, nm in enumerate(names):
        c = ws.cell(row=3 + i, column=1, value=nm)
        c.font = Font(name=FONT)
        c.border = BOX
    fibre_rng = book.data(ws, 3, 2, [2.3, 12.1, 5.4, 3.4, 10.1], fmt="0.0")
    sR_rng = book.data(ws, 3, 3, [0.293, 0.563, 0.390, 0.347, 0.575], fmt="0.000")
    sr_rng = book.data(ws, 3, 4, [0.198, 0.358, 0.264, 0.232, 0.391], fmt="0.000")

    book.text(ws, 9, 1, "觀察：s_R 隨含量上升（用 SLOPE/INTERCEPT 配一條參考線）", bold=True, size=11, color="0F4C81")
    book.header(ws, 10, ["項目", "結果", "公式", "說明"])
    slope_sr = book.calc(ws, 11, "斜率 (s_R ~ 含量)", f"=SLOPE({sR_rng},{fibre_rng})", "含量每增加1%，s_R 大約增加多少", fmt="0.0000")
    intercept_sr = book.calc(ws, 12, "截距", f"=INTERCEPT({sR_rng},{fibre_rng})", "", fmt="0.0000")

    book.text(ws, 14, 1, "乾燥固定項：恆重失重 ±2 mg（矩形）對 1 g 樣品的貢獻", bold=True, size=11, color="0F4C81")
    book.header(ws, 15, ["項目", "結果", "公式", "說明"])
    a_dry = book.calc(ws, 16, "恆重失重半寬 a (mg)", 2, "乾燥至恆重的稱重容許差", fmt="0.00")
    book.mark_input(ws, a_dry)
    sample_g = book.calc(ws, 17, "樣品重 (g)", 1, "約 1 g 樣品", fmt="0.00")
    book.mark_input(ws, sample_g)
    u_dry_calc = book.calc(ws, 18, "u(乾燥) 計算值 (% m/m)", f"=({a_dry}/SQRT(3))/{sample_g}/10",
                            "矩形 a/√3，再換算成占 1 g 樣品的百分比；QUAM 原文採用值 0.115", fmt="0.0000")
    u_dry = book.calc(ws, 19, "u(乾燥)：採用值 (% m/m)", 0.115, "QUAM A6 原文採用值（低含量時不可忽略）", fmt="0.0000")
    book.mark_input(ws, u_dry)

    book.text(ws, 21, 1, "不確定度預算：三個含量層級（分含量層級報告）", bold=True, size=12, color="0F4C81")

    def level_block(start_row, label, sR_val, content_level):
        budget_header(ws, start_row)
        r = start_row + 1
        budget_cell(ws, r, 1, f"s_R（{label}）")
        val_c = budget_cell(ws, r, 2, sR_val, fmt="0.000", fill=INPUT_FILL)
        lvl_label = ws.cell(row=r, column=9, value="纖維含量水平 (%)：")
        lvl_label.font = Font(name=FONT, italic=True, size=9)
        lvl_c = budget_cell(ws, r, 10, content_level, fmt="0.00", fill=INPUT_FILL)
        budget_cell(ws, r, 3, "協同試驗再現性標準差")
        budget_cell(ws, r, 4, "1")
        u_sR = budget_cell(ws, r, 5, f"={val_c.coordinate}", fmt="0.0000")
        rel_sR = budget_cell(ws, r, 6, "—")
        pct_sR = budget_cell(ws, r, 7, None, fmt="0.0")
        r2 = r + 1
        budget_cell(ws, r2, 1, "乾燥固定項")
        budget_cell(ws, r2, 2, f"={u_dry}", fmt="0.0000")
        budget_cell(ws, r2, 3, "矩形（見上方換算）")
        budget_cell(ws, r2, 4, "—")
        u_d = budget_cell(ws, r2, 5, f"={u_dry}", fmt="0.0000")
        rel_d = budget_cell(ws, r2, 6, "—")
        pct_d = budget_cell(ws, r2, 7, None, fmt="0.0")
        r3 = r2 + 1
        uc = budget_cell(ws, r3, 1, "合成 u_c", bold=True)
        uc_v = budget_cell(ws, r3, 5, f"=SQRT(SUMSQ({u_sR.coordinate},{u_d.coordinate}))", fmt="0.0000", bold=True)
        U_v = budget_cell(ws, r3, 6, f"=2*{uc_v.coordinate}", fmt="0.0000", bold=True)
        budget_cell(ws, r3, 3, "U=2×u_c (k=2)")
        pct_v = budget_cell(ws, r3, 7, f"={U_v.coordinate}/{lvl_c.coordinate}*100", fmt="0.0", bold=True)
        budget_cell(ws, r3, 4, "相對 U%→")
        for cell, ucell in [(pct_sR, u_sR), (pct_d, u_d)]:
            cell.value = fx(f"={ucell.coordinate}^2/{uc_v.coordinate}^2*100")
        return val_c, uc_v, U_v, pct_v

    val25, uc25, U25, pct25 = level_block(22, "含量2.5%", 0.292, 2.5)
    val5, uc5, U5, pct5 = level_block(26, "含量5%", 0.400, 5)
    val10, uc10, U10, pct10 = level_block(30, "含量10%", 0.600, 10)

    book.text(ws, 35, 1, "三個層級比較（相對不確定度隨含量遞減）", bold=True, size=11, color="0F4C81")
    book.header(ws, 36, ["含量 (%)", "u_c", "U (k=2)", "相對 U (%)"])
    for i, (v, u, U, name) in enumerate([(2.5, uc25, U25, "2.5%"), (5, uc5, U5, "5%"), (10, uc10, U10, "10%")]):
        row = 37 + i
        budget_cell(ws, row, 1, v, fmt="0.00")
        budget_cell(ws, row, 2, f"={u.coordinate}", fmt="0.0000")
        budget_cell(ws, row, 3, f"={U.coordinate}", fmt="0.0000")
        budget_cell(ws, row, 4, f"={U.coordinate}/{v}*100", fmt="0.0")

    # =================================================================
    # 工作表 2：標準曲線反推濃度的不確定度（QUAM E.4）
    # =================================================================
    ws2 = book.sheet("標準曲線反推", [30, 16, 44, 60])
    book.text(ws2, 1, 1, "校正數據（同第4章 Group A 鈉標準曲線）", bold=True, size=12, color="0F4C81")
    book.header(ws2, 2, ["濃度 x (μg/mL)"])
    x_rng = book.data(ws2, 3, 1, [1.0, 3.0, 5.0, 10.0, 20.0], fmt="0.0")
    book.header(ws2, 2, ["發射訊號 y"], col=2)
    y_rng = book.data(ws2, 3, 2, [0.050, 0.140, 0.242, 0.521, 0.998], fmt="0.000")

    book.header(ws2, 9, ["項目", "結果", "公式", "說明"])
    n = book.calc(ws2, 10, "n：校正點數", f"=COUNT({x_rng})", "有幾個標準品濃度", fmt="0")
    b1 = book.calc(ws2, 11, "b1：斜率", f"=SLOPE({y_rng},{x_rng})", "每單位濃度訊號增加多少", fmt="0.000000")
    b0 = book.calc(ws2, 12, "b0：截距", f"=INTERCEPT({y_rng},{x_rng})", "", fmt="0.000000")
    S = book.calc(ws2, 13, "S：殘差標準差", f"=STEYX({y_rng},{x_rng})", "STEYX 用自由度 n−2，是公式要的 S", fmt="0.000000")
    xbar = book.calc(ws2, 14, "x̄：濃度平均", f"=AVERAGE({x_rng})", "", fmt="0.0000")
    Sxx = book.calc(ws2, 15, "S_xx：Σ(xi−x̄)²", f"=DEVSQ({x_rng})", "DEVSQ 直接算離均差平方和", fmt="0.0000")
    p_reps = book.calc(ws2, 16, "p：未知樣品重複測定次數", 3, "未知樣品測 3 次取平均", fmt="0")
    book.mark_input(ws2, p_reps)
    y_obs = book.calc(ws2, 17, "未知樣品 3 次平均訊號 y_obs", 0.555, "", fmt="0.000")
    book.mark_input(ws2, y_obs)
    x_pred = book.calc(ws2, 18, "x_pred：反推濃度", f"=({y_obs}-{b0})/{b1}", "(y_obs − b0)/b1", fmt="0.0000", key=True)

    book.text(ws2, 20, 1, "var(x_pred) 公式逐項拆開（QUAM E.4）", bold=True, size=12, color="0F4C81")
    book.header(ws2, 21, ["項目", "結果", "公式", "說明"])
    term_p = book.calc(ws2, 22, "1/p", f"=1/{p_reps}", "未知樣品重複測定次數項", fmt="0.0000")
    term_n = book.calc(ws2, 23, "1/n", f"=1/{n}", "校正點數項", fmt="0.0000")
    term_x = book.calc(ws2, 24, "(x_pred−x̄)²/S_xx", f"=({x_pred}-{xbar})^2/{Sxx}", "離校正中心越遠這一項越大", fmt="0.0000")
    sum_terms = book.calc(ws2, 25, "三項合計", f"={term_p}+{term_n}+{term_x}", "", fmt="0.0000")
    factor = book.calc(ws2, 26, "S²/b1²", f"={S}^2/{b1}^2", "殘差變異除以斜率平方", fmt="0.000000")
    var_xpred = book.calc(ws2, 27, "var(x_pred)", f"={factor}*{sum_terms}", "(S²/b1²) × 三項合計", fmt="0.000000", key=True)
    u_xpred = book.calc(ws2, 28, "u(x_pred)", f"=SQRT({var_xpred})", "", fmt="0.0000", key=True)
    U_k2 = book.calc(ws2, 29, "U (k=2)", f"=2*{u_xpred}", "", fmt="0.0000", key=True)
    book.text(ws2, 30, 1, "報告：Na = (11.1 ± 0.4) μg/mL（k=2）", bold=True)

    book.check("粗纖維 2.5% u_c", "粗纖維協同試驗", uc25.coordinate, 0.3138295716)
    book.check("粗纖維 2.5% U", "粗纖維協同試驗", U25.coordinate, 0.6276591432)
    book.check("粗纖維 5% u_c", "粗纖維協同試驗", uc5.coordinate, 0.4162030754)
    book.check("粗纖維 10% u_c", "粗纖維協同試驗", uc10.coordinate, 0.6109214352)
    book.check("粗纖維 10% 相對U%", "粗纖維協同試驗", pct10.coordinate, 12.21842870)
    book.check("標準曲線 斜率 b1", "標準曲線反推", b1, 0.05039948007)
    book.check("標準曲線 殘差標準差 S", "標準曲線反推", S, 0.01380782312)
    book.check("標準曲線 S_xx", "標準曲線反推", Sxx, 230.8)
    book.check("反推濃度 x_pred", "標準曲線反推", x_pred, 11.069875)
    book.check("反推濃度 u(x_pred)", "標準曲線反推", u_xpred, 0.208586234)
    book.check("反推濃度 U (k=2)", "標準曲線反推", U_k2, 0.4171724679)
    return book
