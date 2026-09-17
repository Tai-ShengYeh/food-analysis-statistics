"""Ch04 標準曲線與線性迴歸：OLS 逐步計算 + 1/x² 加權迴歸（WLS）。

R 的答案來源：Rscript R/ch04_regression.R
（鈉標準曲線 Nielsen 章末習題4 Group A；HPLC 咖啡因 5濃度×3重複的加權迴歸範例）。
"""
from excel_common import Book, FONT, BOX
from openpyxl.styles import Font


def build():
    book = Book("ch04_calibration_wls.xlsx",
                 "Ch4 標準曲線與線性迴歸：OLS 逐步計算與 1/x² 加權迴歸（WLS）", "ch04.html")
    book.readme([
        "# 這本活頁簿在做什麼",
        "「標準曲線迴歸」工作表：鈉標準曲線 5 個濃度點（Nielsen 章末習題4 Group A），用 SLOPE/INTERCEPT/RSQ/STEYX 逐步"
        "算出 y=ax+b、r²、殘差標準誤，再把每一點的殘差算出來，最後反推一個未知樣品訊號的濃度。",
        "「WLS_1x2加權」工作表：HPLC 咖啡因 5 個濃度（1~100 ppm）各測 3 次，示範當高濃度訊號的變異明顯比低濃度大"
        "（異方差）時，怎麼用 SUMPRODUCT 公式做 1/x² 加權最小平方法（WLS），並把 OLS 與 WLS 的斜率、截距、"
        "每一點回算濃度的相對誤差 %RE 並排比較。",
        "# 怎麼用",
        "1. 黃色是原始數據（濃度、訊號），其他都是活公式，C 欄印出公式原文對照。",
        "2.「標準曲線迴歸」把 B15 的未知樣品訊號改一改，看反推濃度怎麼變；也可以改 A2:B6 的原始數據看整條線怎麼變。",
        "3.「WLS_1x2加權」看『%RE 回算對照』表：OLS 在 1 ppm（最低濃度）那一列的 %RE 特別大，WLS 1/x² 那一列則小很多——"
        "這就是加權要解決的問題。",
        "4.「對照R答案」工作表核對 Excel 算出來的與 R 的答案是否一致。",
        "# 用 Excel 資料分析工具箱算迴歸（選用，適用於標準曲線）",
        "啟用方式：檔案 → 選項 → 增益集 → 管理『Excel 增益集』→ 勾選『分析工具箱』→ 確定，"
        "『資料』分頁會多一個『資料分析』按鈕。",
        "操作：點『資料分析』→ 選『迴歸』→『輸入Y範圍』框選訊號 y（例如 B2:B6）、『輸入X範圍』框選濃度 x（例如 A2:A6）→"
        "勾選『信賴度』95% → 指定輸出範圍 → 確定。",
        "對照：輸出表『係數』欄的『截距』＝本活頁簿的 INTERCEPT／R 的 coef(fit)[1]；『X』那一列的係數＝SLOPE／"
        "coef(fit)[2]；『R 平方』＝ RSQ／summary(fit)$r.squared；『標準誤』（迴歸統計區塊的那個，不是係數旁的）＝ "
        "STEYX／summary(fit)$sigma；『殘差輸出』勾選後會多印一張逐點殘差表，對應本頁的殘差欄。",
        "⚠️ 工具箱輸出是『靜態的』：改了 A、B 欄的數據，這張輸出表不會自動更新，要重新跑一次『資料分析』。"
        "本活頁簿「標準曲線迴歸」工作表全部用 SLOPE/INTERCEPT/RSQ/STEYX 公式，數據一改就自動重算。"
        "工具箱也沒有『加權迴歸』選項，1/x² 加權一定要靠「WLS_1x2加權」工作表的 SUMPRODUCT 公式或改用 R。",
        "# 本章最重要的觀念",
        "r／r² 只衡量『假設是直線時，貼得多好』，不能證明資料真的是直線——一定要看殘差圖，殘差應隨機散布在 0 附近，"
        "不能有彎曲趨勢。反推濃度一定是 x=(y−b)/a，不要把訊號直接代入 x 去算 y（方向弄反）。",
        "『濃度範圍很寬』本身不是自動加權的理由，要先看：①各濃度重複的 SD 是否隨濃度一路變大；②OLS 殘差是否呈喇叭狀。"
        "選權重不能只看 r²（三個模型的 r² 幾乎一樣），要看低濃度端的回算 %RE 是否明顯改善。",
    ])

    # ---- 標準曲線迴歸（Na 曲線，OLS）------------------------------------
    ws = book.sheet("標準曲線迴歸", [18, 16, 44, 60])
    book.header(ws, 1, ["濃度 x (ug/mL)", "訊號 y"])
    rx = book.data(ws, 2, 1, [1.0, 3.0, 5.0, 10.0, 20.0], fmt="0.0")
    ry = book.data(ws, 2, 2, [0.050, 0.140, 0.242, 0.521, 0.998], fmt="0.000")

    book.header(ws, 8, ["項目", "結果", "公式", "說明"])
    n = book.calc(ws, 9, "n（標準品數）", f"=COUNT({rx})", fmt="0")
    slope = book.calc(ws, 10, "斜率 a (SLOPE)", f"=SLOPE({ry},{rx})", "靈敏度：訊號隨濃度增加的幅度", key=True, fmt="0.00000")
    intercept = book.calc(ws, 11, "截距 b (INTERCEPT)", f"=INTERCEPT({ry},{rx})", "濃度為0時的空白背景訊號", key=True, fmt="0.00000")
    r2 = book.calc(ws, 12, "判定係數 r² (RSQ)", f"=RSQ({ry},{rx})", "y 的變異能被 x 解釋的比例", key=True, fmt="0.0000")
    r = book.calc(ws, 13, "相關係數 r (CORREL)", f"=CORREL({rx},{ry})", "r=√r²；經驗法則 r≥0.997 才理想", fmt="0.0000")
    steyx = book.calc(ws, 14, "殘差標準誤 STEYX", f"=STEYX({ry},{rx})", "殘差的標準差，越小代表點越貼線", fmt="0.00000")
    yobs = book.calc(ws, 15, "未知樣品訊號 y_obs", 0.555, "可以改成你自己的訊號值試試", fmt="0.000")
    book.mark_input(ws, yobs)
    xpred = book.calc(ws, 16, "反推濃度 x_pred", f"=({yobs}-{intercept})/{slope}", "x=(y−b)/a；千萬別把方向弄反！", key=True, fmt="0.0000")

    book.text(ws, 18, 1, "殘差表：逐點檢查『直線』這個假設站不站得住腳（應隨機散布在 0 附近，不能有彎曲趨勢）。", bold=True)
    book.header(ws, 19, ["x", "y實測", "y預測 (=a·x+b)", "殘差 (=y實測−y預測)"])
    x_vals = [1.0, 3.0, 5.0, 10.0, 20.0]
    resid_first = None
    for i in range(5):
        row = 20 + i
        data_row = 2 + i
        c1 = ws.cell(row=row, column=1, value=f"=A{data_row}")
        c1.number_format = "0.0"; c1.font = Font(name=FONT); c1.border = BOX
        c2 = ws.cell(row=row, column=2, value=f"=B{data_row}")
        c2.number_format = "0.000"; c2.font = Font(name=FONT); c2.border = BOX
        c3 = ws.cell(row=row, column=3, value=f"={slope}*A{row}+{intercept}")
        c3.number_format = "0.0000"; c3.font = Font(name=FONT); c3.border = BOX
        c4 = ws.cell(row=row, column=4, value=f"=B{row}-C{row}")
        c4.number_format = "0.0000"; c4.font = Font(name=FONT); c4.border = BOX
        if i == 0:
            resid_first = f"D{row}"

    # ---- WLS 1/x^2 加權迴歸 -----------------------------------------------
    ws2 = book.sheet("WLS_1x2加權", [14, 14, 16, 16, 16, 60])
    book.header(ws2, 1, ["濃度 x (ppm)", "面積 y (訊號)"])
    conc = [1, 1, 1, 5, 5, 5, 10, 10, 10, 50, 50, 50, 100, 100, 100]
    area = [101, 99, 102, 498, 505, 492, 1005, 992, 1018, 4930, 5070, 5005, 9720, 10380, 10040]
    rx2 = book.data(ws2, 2, 1, conc, fmt="0")
    ry2 = book.data(ws2, 2, 2, area, fmt="0")

    book.text(ws2, 18, 1, "加權最小平方法：權重 w=1/x²，直接用 SUMPRODUCT 算出加權平方和，不需要陣列公式。", bold=True)
    book.header(ws2, 19, ["項目", "結果", "公式", "說明"])
    Sw = book.calc(ws2, 20, "Σw （w=1/x²）", f"=SUMPRODUCT(1/{rx2}^2)", fmt="0.000000")
    Swx = book.calc(ws2, 21, "Σw·x", f"=SUMPRODUCT(1/{rx2}^2,{rx2})", fmt="0.0000")
    Swy = book.calc(ws2, 22, "Σw·y", f"=SUMPRODUCT(1/{rx2}^2,{ry2})", fmt="0.0000")
    Swxx = book.calc(ws2, 23, "Σw·x²", f"=SUMPRODUCT(1/{rx2}^2,{rx2},{rx2})", fmt="0.0000")
    Swxy = book.calc(ws2, 24, "Σw·x·y", f"=SUMPRODUCT(1/{rx2}^2,{rx2},{ry2})", fmt="0.0000")
    slope_w = book.calc(ws2, 25, "WLS 斜率 a_w", f"=({Sw}*{Swxy}-{Swx}*{Swy})/({Sw}*{Swxx}-{Swx}^2)",
                         "加權常態方程式解出來的斜率", key=True, fmt="0.0000")
    intercept_w = book.calc(ws2, 26, "WLS 截距 b_w", f"=({Swy}-{slope_w}*{Swx})/{Sw}",
                             key=True, fmt="0.0000")
    slope_ols = book.calc(ws2, 27, "OLS 斜率 a (SLOPE，對照用)", f"=SLOPE({ry2},{rx2})", fmt="0.0000")
    intercept_ols = book.calc(ws2, 28, "OLS 截距 b (INTERCEPT，對照用)", f"=INTERCEPT({ry2},{rx2})", fmt="0.0000")

    book.text(ws2, 30, 1, "%RE 回算對照：把每一筆標準品的訊號當成『未知樣品』反推濃度，看回算的相對誤差 %RE。", bold=True)
    book.header(ws2, 31, ["濃度 x", "面積 y", "OLS回算濃度", "OLS %RE", "WLS回算濃度", "WLS %RE"])
    re_first_ols = None
    re_first_wls = None
    for i in range(len(conc)):
        row = 32 + i
        data_row = 2 + i
        c1 = ws2.cell(row=row, column=1, value=f"=A{data_row}")
        c1.number_format = "0"; c1.font = Font(name=FONT); c1.border = BOX
        c2 = ws2.cell(row=row, column=2, value=f"=B{data_row}")
        c2.number_format = "0"; c2.font = Font(name=FONT); c2.border = BOX
        c3 = ws2.cell(row=row, column=3, value=f"=(B{row}-{intercept_ols})/{slope_ols}")
        c3.number_format = "0.0000"; c3.font = Font(name=FONT); c3.border = BOX
        c4 = ws2.cell(row=row, column=4, value=f"=(C{row}-A{row})/A{row}*100")
        c4.number_format = "0.0000"; c4.font = Font(name=FONT); c4.border = BOX
        c5 = ws2.cell(row=row, column=5, value=f"=(B{row}-{intercept_w})/{slope_w}")
        c5.number_format = "0.0000"; c5.font = Font(name=FONT); c5.border = BOX
        c6 = ws2.cell(row=row, column=6, value=f"=(E{row}-A{row})/A{row}*100")
        c6.number_format = "0.0000"; c6.font = Font(name=FONT); c6.border = BOX
        if i == 0:
            re_first_ols = f"D{row}"
            re_first_wls = f"F{row}"

    last_re_row = 32 + len(conc) - 1
    book.text(ws2, last_re_row + 2, 1,
              "看第一列（1 ppm）：OLS 的 %RE 明顯比 WLS 1/x² 大很多——OLS 為了討好高濃度點，"
              "把低濃度端的反推誤差犧牲掉了。三個模型的 r² 幾乎一樣，看不出這個差別，選權重不能只看 r²。",
              bold=True)

    # ---- 對照 R 答案 -------------------------------------------------------
    book.check("標準曲線斜率 a", "標準曲線迴歸", slope, 0.0503994800693241)
    book.check("標準曲線截距 b", "標準曲線迴歸", intercept, -0.00291594454072781)
    book.check("判定係數 r²", "標準曲線迴歸", r2, 0.999025324016803)
    book.check("相關係數 r", "標準曲線迴歸", r, 0.999512543201336)
    book.check("殘差標準誤 STEYX", "標準曲線迴歸", steyx, 0.013807823116001)
    book.check("反推濃度 x_pred (y=0.555)", "標準曲線迴歸", xpred, 11.0698750021492)
    book.check("第1點殘差", "標準曲線迴歸", resid_first, 0.00251646447140376, tol=1e-6)
    book.check("WLS 1/x² 斜率", "WLS_1x2加權", slope_w, 100.151875454511)
    book.check("WLS 1/x² 截距", "WLS_1x2加權", intercept_w, 0.431545910360954)
    book.check("OLS 斜率（對照）", "WLS_1x2加權", slope_ols, 100.429367890407)
    book.check("OLS 截距（對照）", "WLS_1x2加權", intercept_ols, -3.78834729484022)
    book.check("1ppm OLS %RE", "WLS_1x2加權", re_first_ols, 4.34034336369626, tol=1e-4)
    book.check("1ppm WLS 1/x² %RE", "WLS_1x2加權", re_first_wls, 0.415946913862442, tol=1e-4)
    return book
