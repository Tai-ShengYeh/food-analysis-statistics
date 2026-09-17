"""Ch03 偵測極限（LOD/LOQ）、方法偵測極限（MDL）與品質管制圖（Shewhart／CuSum）。

R 的答案來源：Rscript R/ch03_method_qc.R
（set.seed(42) 產生 20 個空白讀值；set.seed(7) 產生 25 天 QC 值後人為加入漂移與爆表事件）。
"""
from excel_common import Book, FONT, BOX, INPUT_FILL
from openpyxl.styles import Font


def build():
    book = Book("ch03_lod_control_chart.xlsx",
                 "Ch3 方法能力指標與品質管制：LOD/LOQ、MDL、Shewhart 與 CuSum 管制圖", "ch03.html")
    book.readme([
        "# 這本活頁簿在做什麼",
        "「LOD_LOQ」工作表：對空白樣品測 20 次，用空白的平均與標準差算出偵測極限 LOD（X̄+3s）與定量極限 LOQ（X̄+10s）。",
        "「MDL」工作表：用 7 次加標基質樣品的重複，依 EPA 定義算方法偵測極限 MDL = t(n−1,0.99)×SD。",
        "「Shewhart_CuSum」工作表：25 天的 QC 標準品（蛋白質 12.00%）監控數據——這組數據不是隨手打的，"
        "是 R 用 set.seed(7) 產生後，人為加入第 8~11 天的系統性漂移（+0.28%）與第 18 天的單點爆表（12.62%），"
        "目的是讓你同時練到「連續同側」與「單點超界」兩種訊號。管制界限用固定的歷史 s=0.15，並用 IF 公式自動標記"
        "哪些天超出 ±2s 警告界限、哪些天超出 ±3s 行動界限；CuSum 欄用累加公式，示範它比 Shewhart 更早抓到微小漂移。",
        "# 怎麼用",
        "1. 黃色是原始數據（空白讀值、加標讀值、QC 值），其他都是活公式，C 欄印出公式原文對照。",
        "2.「Shewhart_CuSum」把 B2 的目標值、B3 的歷史 s 改一改，看警告線／行動線怎麼跟著平移或變寬變窄。",
        "3. 試著把某一天的 QC 值改到超過 12.45 或低於 11.55，看 D 欄的判定欄位會不會自動跳出「超出±3s」。",
        "4.「對照R答案」工作表核對 Excel 算出來的與 R 的答案是否一致。",
        "# 用 Excel 資料分析工具箱算敘述統計（選用）",
        "檔案 → 選項 → 增益集 → 管理『Excel 增益集』→ 勾選『分析工具箱』→ 確定，『資料』分頁會多一個『資料分析』按鈕。",
        "點『資料分析』→『敘述統計』→ 輸入範圍框選空白 20 個讀值 → 勾選『摘要統計』→ 確定，輸出表的『平均值』"
        "與『標準差』就對應 LOD/LOQ 公式裡的 X̄_Blk 與 s_Blk，你還是要自己在旁邊多打一格 =平均值+3*標準差 才算出 LOD。",
        "⚠️ 工具箱輸出是『靜態的』：改了空白讀值，這張輸出表不會自動更新，要重新跑一次。"
        "本活頁簿「LOD_LOQ」工作表全部是公式，數據一改就自動重算。",
        "# 本章最重要的觀念",
        "管制界限（±2s／±3s）一定要用『方法平常本來的變異』——也就是方法確效或歷史穩定期得到的 s（本例固定 s=0.15）。"
        "千萬不要拿『正在監控的這一批』QC 值自己算 STDEV 當管制界限：一旦這批資料已經出現漂移或爆表，"
        "異常點會把 STDEV 一起拉大，界限跟著變寬，反而更抓不到問題——等於拿問題本身去稀釋判斷問題的基準。"
        "這就是為什麼「Shewhart_CuSum」工作表的歷史 s 是一個獨立的黃色輸入格，不是 =STDEV(QC值)。",
        "另一個重點：管制圖不是只看『有沒有點超出 ±3s』——連續多點落在中心線同側（本例第 7~16 天）"
        "本身就是系統性偏移的訊號，CuSum 欄的累加曲線會比 Shewhart 更早、更明顯地把這種小漂移顯示出來。",
    ])

    # ---- LOD / LOQ -----------------------------------------------------
    ws = book.sheet("LOD_LOQ", [18, 16, 44, 60])
    book.header(ws, 1, ["空白讀值 (mg/L)"])
    blank_vals = [0.011, 0.0068, 0.0088, 0.0094, 0.0089, 0.0078, 0.0113, 0.0078, 0.0124, 0.0079,
                  0.0109, 0.013, 0.0049, 0.0074, 0.0077, 0.0094, 0.0074, 0.0022, 0.0026, 0.0109]
    rng = book.data(ws, 2, 1, blank_vals, fmt="0.0000")
    book.header(ws, 23, ["項目", "結果", "公式", "說明"])
    n = book.calc(ws, 24, "n（空白測定次數）", f"=COUNT({rng})", fmt="0")
    x_blk = book.calc(ws, 25, "空白平均 X̄Blk", f"=AVERAGE({rng})", "X̄_Blk", fmt="0.0000")
    s_blk = book.calc(ws, 26, "空白標準差 sBlk", f"=STDEV.S({rng})", "s_Blk（n−1）", fmt="0.0000")
    lod = book.calc(ws, 27, "偵測極限 LOD", f"={x_blk}+3*{s_blk}", "X̄Blk + 3×sBlk（Nielsen式4.19）", key=True, fmt="0.0000")
    loq = book.calc(ws, 28, "定量極限 LOQ", f"={x_blk}+10*{s_blk}", "X̄Blk + 10×sBlk", key=True, fmt="0.0000")
    book.text(ws, 30, 1,
              "解讀：低於 LOD → 報「未檢出 ND」；LOD~LOQ 之間只能定性、數字不可靠；高於 LOQ 才能安心報定量值。",
              bold=True)

    # ---- MDL -------------------------------------------------------------
    ws2 = book.sheet("MDL", [18, 16, 44, 60])
    book.header(ws2, 1, ["加標基質讀值 (mg/L)"])
    rng2 = book.data(ws2, 2, 1, [1.02, 0.98, 1.05, 0.99, 1.01, 0.97, 1.03], fmt="0.00")
    book.header(ws2, 10, ["項目", "結果", "公式", "說明"])
    n2 = book.calc(ws2, 11, "n（加標重複次數）", f"=COUNT({rng2})", "EPA 定義要求 n≥7", fmt="0")
    mean2 = book.calc(ws2, 12, "加標平均", f"=AVERAGE({rng2})", fmt="0.0000")
    sd2 = book.calc(ws2, 13, "加標標準差 SD", f"=STDEV.S({rng2})", fmt="0.0000")
    t99 = book.calc(ws2, 14, "t(n−1, 0.99) 單尾", f"=T.INV(0.99,{n2}-1)", "R 的 qt(0.99,6) 是單尾 99% 分位數，"
                     "Excel 要用單尾的 T.INV(機率,自由度)，不是雙尾的 T.INV.2T！", fmt="0.00000")
    mdl = book.calc(ws2, 15, "方法偵測極限 MDL", f"={t99}*{sd2}", "t(n-1,0.99) × SD，涵蓋整個方法流程的變異", key=True, fmt="0.0000")
    book.text(ws2, 17, 1, "MDL 用『含分析物的真實基質樣品』重複分析，比只看空白雜訊的 LOD 更嚴謹，涵蓋萃取、淨化、上機整個流程的變異。", bold=True)

    # ---- Shewhart 與 CuSum 管制圖 -----------------------------------------
    ws3 = book.sheet("Shewhart_CuSum", [10, 14, 16, 30, 50])
    book.header(ws3, 1, ["項目", "結果", "公式", "說明"])
    target = book.calc(ws3, 2, "目標值（中心線 CL）", 12.00, "QC 標準品的已知蛋白質% 標示值", fmt="0.00")
    book.mark_input(ws3, target)
    hist_s = book.calc(ws3, 3, "歷史 s（方法確效期得到）", 0.15, "固定值！不要用本批 QC 的 STDEV，否則異常點會把界限一起拉寬", fmt="0.00")
    book.mark_input(ws3, hist_s)
    uwl = book.calc(ws3, 4, "上警告界限 UWL (+2s)", f"={target}+2*{hist_s}", fmt="0.000")
    lwl = book.calc(ws3, 5, "下警告界限 LWL (−2s)", f"={target}-2*{hist_s}", fmt="0.000")
    ual = book.calc(ws3, 6, "上行動界限 UAL (+3s)", f"={target}+3*{hist_s}", fmt="0.000")
    lal = book.calc(ws3, 7, "下行動界限 LAL (−3s)", f"={target}-3*{hist_s}", fmt="0.000")

    qc_values = [12.343, 11.82, 11.896, 11.938, 11.854, 11.858, 12.112, 12.262, 12.303, 12.608,
                 12.334, 12.408, 12.342, 12.049, 12.284, 12.07, 11.866, 12.62, 11.999, 12.148,
                 12.126, 12.106, 12.196, 11.792, 12.191]
    head_row = 9
    book.header(ws3, head_row, ["分析日", "QC值(%)", "CuSum累積偏差", "判定"])
    first = head_row + 1
    for i, v in enumerate(qc_values):
        row = first + i
        day_cell = ws3.cell(row=row, column=1, value=i + 1)
        day_cell.font = Font(name=FONT)
        day_cell.border = BOX
        val_cell = ws3.cell(row=row, column=2, value=v)
        val_cell.fill = INPUT_FILL
        val_cell.font = Font(name=FONT)
        val_cell.border = BOX
        val_cell.number_format = "0.000"
        if i == 0:
            cusum_formula = f"=B{row}-{target}"
        else:
            cusum_formula = f"=C{row-1}+(B{row}-{target})"
        c_cell = ws3.cell(row=row, column=3, value=cusum_formula)
        c_cell.font = Font(name=FONT)
        c_cell.border = BOX
        c_cell.number_format = "0.000"
        judge_formula = (f'=IF(OR(B{row}>{ual},B{row}<{lal}),"超出±3s(行動)",'
                          f'IF(OR(B{row}>{uwl},B{row}<{lwl}),"超出±2s(警告)","正常"))')
        d_cell = ws3.cell(row=row, column=4, value=judge_formula)
        d_cell.font = Font(name=FONT)
        d_cell.border = BOX
        d_cell.alignment = None

    last = first + len(qc_values) - 1
    book.text(ws3, last + 2, 1,
              "判讀口訣：① 任一點超出±3s→立即停線找根本原因；② 連續7點同側或連續升降→系統性漂移（即使沒超界）；"
              "③ 只超±2s未超±3s→提高警覺、加測確認。CuSum 欄從第8天開始明顯轉正並持續攀升——"
              "比 Shewhart 的判定欄（要到第10天才出現「超出±3s」）更早顯示系統性漂移。", bold=True)

    cusum_day8 = f"C{first + 7}"
    cusum_day25 = f"C{last}"
    qc_day10 = f"B{first + 9}"

    # ---- 對照 R 答案 -------------------------------------------------------
    book.check("空白平均 X̄Blk", "LOD_LOQ", x_blk, 0.008425)
    book.check("空白標準差 sBlk", "LOD_LOQ", s_blk, 0.0028798528836694)
    book.check("偵測極限 LOD", "LOD_LOQ", lod, 0.0170645586510082)
    book.check("定量極限 LOQ", "LOD_LOQ", loq, 0.037223528836694)
    book.check("方法偵測極限 MDL", "MDL", mdl, 0.0902011269039659)
    book.check("t(6,0.99) 單尾", "MDL", t99, 3.14266840329098)
    book.check("CuSum 第8天累積偏差", "Shewhart_CuSum", cusum_day8, 0.0830000000000002, tol=1e-6)
    book.check("CuSum 第25天累積偏差", "Shewhart_CuSum", cusum_day25, 3.525, tol=1e-6)
    book.check("QC 第10天讀值（資料核對）", "Shewhart_CuSum", qc_day10, 12.608, tol=1e-9)
    return book
