"""Ch05 有效數字與異常值檢定：Dixon Q 檢定、Grubbs 檢定。

R 的答案來源：Rscript R/ch05_outliers_sigfig.R
（水分含異常值 55.31 的 5 筆版本；乾物質習題 4 筆版本；grubbs_crit() 的 t 分布近似公式）。
"""
from excel_common import Book, FONT, BOX, INPUT_FILL
from openpyxl.styles import Font


def build():
    book = Book("ch05_sigfig_outliers.xlsx",
                 "Ch5 有效數字取位與異常值檢定：Dixon Q 與 Grubbs 檢定", "ch05.html")
    book.readme([
        "# 這本活頁簿在做什麼",
        "「有效數字取位」工作表：示範乘除看『有效位數最少』、加減看『小數位數最少』的取位規則，"
        "以及課本經典的稀釋倍數陷阱（50 當成 1 位有效數字 vs 正確的 50.0）。",
        "「DixonQ檢定」工作表：用 Q=gap/全距 判斷可疑值能不能捨棄，臨界值查表用 INDEX/MATCH 依 n 自動取值"
        "（不用每次手動翻書）。示範水分 5 筆（含可疑值 55.31）與乾物質習題 4 筆兩個案例。",
        "「Grubbs檢定」工作表：用 G=|可疑值−平均|/SD 判斷，臨界值查表用 VLOOKUP 依 n 取值——"
        "這張查表不是查課本（Grubbs 臨界值課本沒有列），而是用 R 的 grubbs_crit() 公式（t 分布近似）"
        "實際算出來後填進表裡，比 Excel 內建函數更穩定、也不用每次重新設定公式。",
        "# 怎麼用",
        "1. 黃色是原始數據，其他都是活公式，C 欄印出公式原文對照。",
        "2.「DixonQ檢定」「Grubbs檢定」都把可疑值改一改（例如把 55.31 改成 60），看判定欄會不會從「可剔除」變回「保留」。",
        "3.「對照R答案」工作表核對 Excel 算出來的與 R 的答案是否一致。",
        "# 資料分析工具箱：這一章沒有直接對應的功能",
        "Dixon Q 檢定與 Grubbs 檢定都不是 Excel『資料分析工具箱』裡的項目（工具箱只有敘述統計、t/F 檢定、ANOVA、迴歸等）。"
        "工具箱裡唯一用得上的是『敘述統計』，可以幫你快速算出平均值、標準差，但查臨界值、算 Q 或 G、下判定，"
        "還是要靠本活頁簿的 INDEX/MATCH、VLOOKUP 與 IF 公式（或直接用 R）。",
        "# 本章最重要的觀念",
        "有效數字規則不能機械套用：稀釋倍數『50』如果是用 A 級容量瓶配的，它其實是 50.0（3 位有效數字），"
        "當成 1 位會嚴重低估最終濃度（2175 被誤算成 2000，正確應該是 2180）。",
        "Q 檢定、Grubbs 檢定拒絕 H0（判定為異常值）不等於可以直接刪數據——必須同時有『可追溯的操作失誤紀錄』"
        "與『統計檢定支持』，兩者兼備才可刪除，且要在報告中揭露；純粹用統計理由拒絕數值通常不明智（QUAM §2.4.13）。",
    ])

    # ---- 有效數字取位 ---------------------------------------------------
    ws = book.sheet("有效數字取位", [34, 16, 50, 60])
    book.header(ws, 1, ["項目", "結果", "公式", "說明"])
    ex1 = book.calc(ws, 2, "36.54×238×1.1 取有效數字", "=ROUND(9566.172,2-1-INT(LOG10(ABS(9566.172))))",
                     "9566.172 中 1.1 只有 2 位有效數字，乘除看最少的那一項", key=True, fmt="0")
    ex2 = book.calc(ws, 3, "0.01672 取3位有效數字", "=ROUND(0.01672,3-1-INT(LOG10(ABS(0.01672))))",
                     "前導零不算有效位數", fmt="0.0000")
    ex3 = book.calc(ws, 4, "7.45+8.725 (加法看小數位)", "=ROUND(7.45+8.725,2)",
                     "7.45 只有 2 位小數，加減法取位看小數位數最少的一項", key=True, fmt="0.00")
    ex4 = book.calc(ws, 5, "咖啡因43.5稀釋『50』(誤當1位)", "=ROUND(43.5*50,1-1-INT(LOG10(ABS(43.5*50))))",
                     "把稀釋倍數50當成1位有效數字 → 嚴重低估！", fmt="0")
    ex5 = book.calc(ws, 6, "咖啡因43.5稀釋『50.0』(正確3位)", "=ROUND(43.5*50,3-1-INT(LOG10(ABS(43.5*50))))",
                     "A級容量瓶配製，50.0才是正確的有效數字位數", key=True, fmt="0")
    book.text(ws, 8, 1, "取位公式 =ROUND(數值, 位數−1−INT(LOG10(ABS(數值)))) 可以直接指定『要留幾位有效數字』，"
              "不用自己數小數點在哪。務必先算完全部步驟，最後才四捨五入一次。", bold=True)

    # ---- Dixon Q 檢定 ------------------------------------------------------
    ws2 = book.sheet("DixonQ檢定", [22, 16, 44, 60])
    book.header(ws2, 1, ["n", "Q0.90 臨界值"])
    q_ns = [3, 4, 5, 6, 7, 8, 9, 10]
    q_crits = [0.94, 0.76, 0.64, 0.56, 0.51, 0.47, 0.44, 0.41]
    for i, (nv, qv) in enumerate(zip(q_ns, q_crits)):
        row = 2 + i
        c1 = ws2.cell(row=row, column=1, value=nv); c1.font = Font(name=FONT); c1.border = BOX; c1.number_format = "0"
        c2 = ws2.cell(row=row, column=2, value=qv); c2.font = Font(name=FONT); c2.border = BOX; c2.number_format = "0.00"
    q_n_range = "$A$2:$A$9"
    q_val_range = "$B$2:$B$9"
    book.text(ws2, 10, 1, "上表是 Nielsen Table 4.4 的 90% 信賴 Dixon Q 臨界值（查課本表格得到，非計算值）。", size=10, color="64748B")

    book.text(ws2, 12, 1, "案例1：水分 5 筆（含可疑低值 55.31）", bold=True)
    book.header(ws2, 13, ["排序後數值 (%)"])
    rngb = book.data(ws2, 14, 1, [55.31, 64.45, 64.53, 64.78, 65.10], fmt="0.00")
    book.header(ws2, 20, ["項目", "結果", "公式", "說明"])
    nb = book.calc(ws2, 21, "n", f"=COUNT({rngb})", fmt="0")
    gap_low_b = book.calc(ws2, 22, "gap_低值端 = x2−x1", "=A15-A14", "最小值可疑時的差距")
    gap_high_b = book.calc(ws2, 23, "gap_高值端 = xn−xn-1", "=A18-A17", "最大值可疑時的差距")
    range_b = book.calc(ws2, 24, "全距 W = xn−x1", "=A18-A14", fmt="0.00")
    Qb = book.calc(ws2, 25, "Q值 = MAX(兩端gap)/W", f"=MAX({gap_low_b},{gap_high_b})/{range_b}", "取兩端較大的那個gap", key=True)
    Qcritb = book.calc(ws2, 26, "查表臨界值 Q0.90", f"=INDEX({q_val_range},MATCH({nb},{q_n_range},0))", "依n用INDEX/MATCH查表", key=True, fmt="0.00")
    judgeb = book.calc(ws2, 27, "判定", f'=IF({Qb}>{Qcritb},"可剔除(異常值)","保留")', fmt="General")

    book.text(ws2, 29, 1, "案例2：乾物質習題 4 筆（可疑低值 82.20）", bold=True)
    book.header(ws2, 30, ["排序後數值 (%)"])
    rngd = book.data(ws2, 31, 1, [82.20, 88.62, 88.74, 89.20], fmt="0.00")
    book.header(ws2, 36, ["項目", "結果", "公式", "說明"])
    nd = book.calc(ws2, 37, "n", f"=COUNT({rngd})", fmt="0")
    gap_low_d = book.calc(ws2, 38, "gap_低值端", "=A32-A31")
    gap_high_d = book.calc(ws2, 39, "gap_高值端", "=A34-A33")
    range_d = book.calc(ws2, 40, "全距 W", "=A34-A31", fmt="0.00")
    Qd = book.calc(ws2, 41, "Q值", f"=MAX({gap_low_d},{gap_high_d})/{range_d}", key=True)
    Qcritd = book.calc(ws2, 42, "查表臨界值 Q0.90", f"=INDEX({q_val_range},MATCH({nd},{q_n_range},0))", key=True, fmt="0.00")
    judged = book.calc(ws2, 43, "判定", f'=IF({Qd}>{Qcritd},"可剔除(異常值)","保留")', fmt="General")

    # ---- Grubbs 檢定 --------------------------------------------------------
    ws3 = book.sheet("Grubbs檢定", [30, 16, 50, 60])
    book.header(ws3, 1, ["n", "Gcrit (α=0.05，R的grubbs_crit()實算)"])
    g_ns = [3, 4, 5, 6, 7, 8, 9, 10]
    g_crits = [1.15430485134404, 1.48125, 1.71503731234336, 1.88714511778393,
               2.0199685076796, 2.12664508719546, 2.21500422332553, 2.2899540844796]
    for i, (nv, gv) in enumerate(zip(g_ns, g_crits)):
        row = 2 + i
        c1 = ws3.cell(row=row, column=1, value=nv); c1.font = Font(name=FONT); c1.border = BOX; c1.number_format = "0"
        c2 = ws3.cell(row=row, column=2, value=gv); c2.font = Font(name=FONT); c2.border = BOX; c2.number_format = "0.00000"
    g_table_range = "$A$2:$B$9"
    book.text(ws3, 10, 1,
              "上表臨界值不是查課本（Grubbs 臨界值 Nielsen 沒有附表），而是用該章 R 檔的 grubbs_crit(n) "
              "（t 分布近似公式）實際算出來後填進去，比每次重打 Excel 公式更不容易出錯。", size=10, color="64748B")

    book.text(ws3, 12, 1, "乾物質習題 4 筆（原始順序，不用排序）", bold=True)
    book.header(ws3, 13, ["乾物質數值 (%)", "|xi−平均| 偏差"])
    dry_vals = [88.62, 88.74, 89.20, 82.20]
    for i, v in enumerate(dry_vals):
        row = 14 + i
        c1 = ws3.cell(row=row, column=1, value=v)
        c1.fill = INPUT_FILL
        c1.font = Font(name=FONT); c1.border = BOX; c1.number_format = "0.00"
        c2 = ws3.cell(row=row, column=2, value=f"=ABS(A{row}-$B$20)")
        c2.font = Font(name=FONT); c2.border = BOX; c2.number_format = "0.0000"
    dev_range = "B14:B17"
    dry_range = "A14:A17"

    book.header(ws3, 19, ["項目", "結果", "公式", "說明"])
    meang = book.calc(ws3, 20, "平均數 x̄", f"=AVERAGE({dry_range})", fmt="0.00")
    sdg = book.calc(ws3, 21, "標準差 SD", f"=STDEV.S({dry_range})", fmt="0.00")
    ng = book.calc(ws3, 22, "n", f"=COUNT({dry_range})", fmt="0")
    Gg = book.calc(ws3, 23, "G值 = MAX(|偏差|)/SD", f"=MAX({dev_range})/{sdg}", key=True)
    Gcritg = book.calc(ws3, 24, "查表臨界值 Gcrit", f"=VLOOKUP({ng},{g_table_range},2,FALSE)", "依n用VLOOKUP查表", key=True, fmt="0.00000")
    judgeg = book.calc(ws3, 25, "判定", f'=IF({Gg}>{Gcritg},"可剔除(異常值，但差距極小)","保留")', fmt="General")
    dry_clean_mean = book.calc(ws3, 26, "捨棄後平均（用LARGE取最大3筆）", f"=AVERAGE(LARGE({dry_range},1),LARGE({dry_range},2),LARGE({dry_range},3))",
                                "假設可疑值是最小值時的取巧寫法：LARGE取第1~3大", key=True, fmt="0.00")
    dry_clean_sd = book.calc(ws3, 27, "捨棄後標準差", f"=STDEV.S(LARGE({dry_range},1),LARGE({dry_range},2),LARGE({dry_range},3))",
                              key=True, fmt="0.00")

    gcrit_formula_check = book.calc(
        ws3, 29, "用公式驗算 Gcrit(n=4)（不查表）",
        f"=(({ng}-1)/SQRT({ng}))*SQRT((T.INV.2T(0.05/{ng},{ng}-2)^2)/({ng}-2+T.INV.2T(0.05/{ng},{ng}-2)^2))",
        "應該和查表值 1.48125 一致；Excel 不查表也能直接用 T.INV.2T 公式重現同一個數字", fmt="0.00000")

    # ---- 對照 R 答案 -------------------------------------------------------
    book.check("有效數字例1 (2位)", "有效數字取位", ex1, 9600, tol=1e-6)
    book.check("有效數字例3 (加法取位)", "有效數字取位", ex3, 16.18, tol=1e-6)
    book.check("稀釋誤當1位 (錯誤示範)", "有效數字取位", ex4, 2000, tol=1e-6)
    book.check("稀釋正確3位", "有效數字取位", ex5, 2180, tol=1e-6)
    book.check("水分5筆 Q值", "DixonQ檢定", Qb, 0.933605720122575)
    book.check("水分5筆 Qcrit(n=5)", "DixonQ檢定", Qcritb, 0.64, tol=1e-6)
    book.check("乾物質4筆 Q值", "DixonQ檢定", Qd, 0.917142857142857)
    book.check("乾物質4筆 Qcrit(n=4)", "DixonQ檢定", Qcritd, 0.76, tol=1e-6)
    book.check("Grubbs G值", "Grubbs檢定", Gg, 1.49578292784366)
    book.check("Grubbs Gcrit(n=4,查表)", "Grubbs檢定", Gcritg, 1.48125)
    book.check("Grubbs Gcrit(n=4,公式驗算)", "Grubbs檢定", gcrit_formula_check, 1.48125)
    book.check("捨棄後平均", "Grubbs檢定", dry_clean_mean, 88.8533333333333)
    book.check("捨棄後標準差", "Grubbs檢定", dry_clean_sd, 0.306159000085468)
    return book
