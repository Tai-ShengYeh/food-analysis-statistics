"""Excel 活頁簿共用工具：給沒有程式基礎的學生「打開就能算」的版本。

每章一個模組 scripts/excel_books/chNN.py，實作 build(book)；由 scripts/build_excel.py 統一產生
downloads/excel/chNN_*.xlsx，並用 LibreOffice 無頭重算、逐格核對「Excel 算出來的值」與「R 的答案」。

每本活頁簿的固定結構：
  1.「說明」      這本在算什麼、對應哪一章、怎麼用、（需要時）資料分析工具箱的操作步驟
  2. 一到數張計算表 黃底＝學生可以改的原始數據；白底＝活公式；右邊欄位寫白話說明
  3.「對照R答案」  每個關鍵結果一列：Excel 值（連結到計算表）、R 的答案、是否一致（✓／✗）

設計原則：公式一律用學生在 Excel 裡真的會打的函數（AVERAGE、STDEV.S、T.INV.2T、SLOPE…），
不要用陣列公式或 LET/LAMBDA 等新函數，Excel 2016 與 LibreOffice 都要能開。
"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT = "Microsoft JhengHei"
INPUT_FILL = PatternFill("solid", fgColor="FFF4C2")    # 黃：可以改的數據
HEAD_FILL = PatternFill("solid", fgColor="0F4C81")     # 深藍：表頭
RESULT_FILL = PatternFill("solid", fgColor="E3F5E5")   # 淡綠：關鍵結果
NOTE_FONT = Font(name=FONT, size=10, color="64748B")
THIN = Side(style="thin", color="CBD5E1")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# Excel 2010 之後才有的函數，寫進檔案時要加 _xlfn. 前綴，否則 Excel 開啟會顯示 #NAME?
_XLFN = ["STDEV.S", "STDEV.P", "VAR.S", "VAR.P", "T.INV.2T", "T.INV", "T.DIST.2T", "T.DIST.RT", "T.DIST",
         "T.TEST", "F.TEST", "F.DIST.RT", "F.DIST", "F.INV.RT", "F.INV", "NORM.S.INV", "NORM.S.DIST",
         "NORM.INV", "NORM.DIST", "CHISQ.INV.RT", "CHISQ.DIST.RT", "CONFIDENCE.T", "CONFIDENCE.NORM",
         "PERCENTILE.INC", "PERCENTILE.EXC", "RANK.EQ", "RANK.AVG"]
_XLFN.sort(key=len, reverse=True)


def fx(formula):
    """把學生看到的公式（=T.INV.2T(0.05,3)）轉成檔案裡要存的形式。"""
    out = formula
    for name in _XLFN:
        out = out.replace(name + "(", "§§" + name + "(")      # 先標記，避免 T.INV 吃到 T.INV.2T
    for name in _XLFN:
        out = out.replace("§§" + name + "(", "_xlfn." + name + "(")
    return out


class Book:
    def __init__(self, filename, title, chapter_html):
        self.filename = filename
        self.title = title
        self.chapter_html = chapter_html
        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        self.checks = []          # (說明, "工作表!A1", R 的答案, 容許差)
        self._readme = None

    # ---- 說明頁 -------------------------------------------------------
    def readme(self, lines):
        """lines：字串清單；以 '# ' 開頭的當小標。"""
        ws = self.wb.create_sheet("說明", 0)
        ws.column_dimensions["A"].width = 110
        ws["A1"] = self.title
        ws["A1"].font = Font(name=FONT, size=16, bold=True, color="0F4C81")
        ws["A2"] = f"對應網頁：{self.chapter_html}　｜　黃色儲存格＝可以改成你自己的數據；其他是公式，請不要直接覆蓋。"
        ws["A2"].font = NOTE_FONT
        row = 4
        for line in lines:
            c = ws.cell(row=row, column=1, value=line[2:] if line.startswith("# ") else line)
            c.font = Font(name=FONT, size=12, bold=True, color="0F4C81") if line.startswith("# ") else Font(name=FONT, size=11)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            row += 1
        self._readme = ws
        return ws

    # ---- 計算表 -------------------------------------------------------
    def sheet(self, name, widths=None):
        ws = self.wb.create_sheet(name)
        for i, w in enumerate(widths or [16, 16, 16, 16, 16, 60], start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        return ws

    def header(self, ws, row, labels, col=1):
        for i, text in enumerate(labels):
            c = ws.cell(row=row, column=col + i, value=text)
            c.font = Font(name=FONT, bold=True, color="FFFFFF")
            c.fill = HEAD_FILL
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = BOX

    def data(self, ws, row, col, values, vertical=True, fmt=None):
        """寫入一串原始數據（黃底），回傳範圍字串，例如 'A2:A5'。"""
        for i, v in enumerate(values):
            r, c = (row + i, col) if vertical else (row, col + i)
            cell = ws.cell(row=r, column=c, value=v)
            cell.fill = INPUT_FILL
            cell.border = BOX
            cell.font = Font(name=FONT)
            if fmt:
                cell.number_format = fmt
        last = (row + len(values) - 1, col) if vertical else (row, col + len(values) - 1)
        return f"{get_column_letter(col)}{row}:{get_column_letter(last[1])}{last[0]}"

    def calc(self, ws, row, label, formula, note="", fmt="0.0000", key=False, label_col=1):
        """一列計算：[標籤][公式][公式文字][白話說明]。回傳公式所在儲存格位址（如 'B7'）。"""
        a = ws.cell(row=row, column=label_col, value=label)
        a.font = Font(name=FONT, bold=key)
        a.border = BOX
        b = ws.cell(row=row, column=label_col + 1, value=fx(formula) if isinstance(formula, str) else formula)
        b.number_format = fmt
        b.border = BOX
        b.font = Font(name=FONT, bold=key)
        if key:
            b.fill = RESULT_FILL
        if isinstance(formula, str) and formula.startswith("="):
            t = ws.cell(row=row, column=label_col + 2, value="'" + formula)     # 顯示公式文字給學生看
            t.font = Font(name="Consolas", size=10, color="1565C0")
        if note:
            n = ws.cell(row=row, column=label_col + 3, value=note)
            n.font = NOTE_FONT
            n.alignment = Alignment(wrap_text=True, vertical="top")
        return f"{get_column_letter(label_col + 1)}{row}"

    def mark_input(self, ws, cell):
        """把某個計算列的儲存格標成黃色（學生可以改的參數，例如信心水準）。"""
        ws[cell].fill = INPUT_FILL

    def text(self, ws, row, col, value, bold=False, size=11, color=None):
        c = ws.cell(row=row, column=col, value=value)
        c.font = Font(name=FONT, bold=bold, size=size, color=color)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        return c

    # ---- 與 R 的答案對照 ----------------------------------------------
    def check(self, label, sheet_name, cell, r_value, tol=None):
        """登記一個要核對的結果。r_value 必須是 Rscript 實際算出來的值。"""
        if tol is None:
            tol = max(abs(r_value) * 5e-4, 1e-9)
        self.checks.append((label, sheet_name, cell, r_value, tol))

    def finish(self, path):
        ws = self.wb.create_sheet("對照R答案")
        for i, w in enumerate([44, 18, 18, 12, 50], start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        self.header(ws, 1, ["項目", "Excel 算出來", "R 的答案", "一致？", "來源儲存格"])
        for i, (label, sheet_name, cell, r_value, tol) in enumerate(self.checks, start=2):
            ref = f"'{sheet_name}'!{cell}"
            ws.cell(row=i, column=1, value=label).font = Font(name=FONT)
            ws.cell(row=i, column=2, value=f"={ref}").number_format = "0.0000"
            ws.cell(row=i, column=3, value=r_value).number_format = "0.0000"
            ws.cell(row=i, column=4, value=f'=IF(ABS(B{i}-C{i})<={tol},"✓","✗")').alignment = Alignment(horizontal="center")
            ws.cell(row=i, column=5, value=ref).font = NOTE_FONT
            for c in range(1, 6):
                ws.cell(row=i, column=c).border = BOX
        n = len(self.checks) + 3
        self.text(ws, n, 1, "如果你改了黃色儲存格的數據，這裡當然會變成 ✗——那是因為 R 的答案是用原始教材數據算的。", color="64748B", size=10)
        self.wb.save(path)
