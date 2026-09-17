#!/usr/bin/env python3
"""產生並驗證各章 Excel 活頁簿（downloads/excel/*.xlsx）。

    python scripts/build_excel.py            # 產生全部 + 用 LibreOffice 重算驗證
    python scripts/build_excel.py ch02 ch16  # 只做指定章
    python scripts/build_excel.py --no-verify

驗證方式：LibreOffice 無頭開檔重算後另存，再讀出每本「對照R答案」工作表的「一致？」欄，
任何 ✗ 或任何工作表出現 #NAME? / #DIV/0! / #VALUE! 等錯誤值都視為失敗（exit 1）。
需要 openpyxl；驗證需要 LibreOffice（soffice）。
"""
import importlib
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "downloads" / "excel"
sys.path.insert(0, str(Path(__file__).resolve().parent))

SOFFICE = [shutil.which("soffice"), r"C:\Program Files\LibreOffice\program\soffice.exe",
           "/usr/bin/soffice", "/Applications/LibreOffice.app/Contents/MacOS/soffice"]


def find_soffice():
    for p in SOFFICE:
        if p and Path(p).exists():
            return p
    return None


def verify(paths):
    soffice = find_soffice()
    if not soffice:
        print("⚠ 找不到 LibreOffice，略過重算驗證。")
        return True
    ok = True
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run([soffice, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", tmp] +
                       [str(p) for p in paths], check=True, capture_output=True, timeout=600)
        for p in paths:
            wb = load_workbook(Path(tmp) / p.name, data_only=True)
            bad = []
            for ws in wb.worksheets:
                for row in ws.iter_rows():
                    for c in row:
                        if isinstance(c.value, str) and c.value.startswith("#") and c.value.endswith(("!", "?", "A")):
                            bad.append(f"{ws.title}!{c.coordinate}={c.value}")
            ws = wb["對照R答案"]
            rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] and r[3] in ("✓", "✗")]
            fails = [f"{r[0]}：Excel {r[1]} vs R {r[2]}" for r in rows if r[3] != "✓"]
            status = "✓" if not bad and not fails and rows else "✗"
            print(f"{status} {p.name}  核對 {len(rows)} 項" + ("" if rows else "（沒有任何核對項！）"))
            for m in bad + fails:
                print("    " + m)
            ok = ok and status == "✓"
    return ok


def main():
    for stream in (sys.stdout, sys.stderr):      # Windows 主控台預設 cp950，印不出 ✓
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8", errors="replace")
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    OUT.mkdir(parents=True, exist_ok=True)
    mods = sorted(p.stem for p in (Path(__file__).parent / "excel_books").glob("ch*.py"))
    if args:
        mods = [m for m in mods if m in args]
    built = []
    for name in mods:
        mod = importlib.import_module(f"excel_books.{name}")
        book = mod.build()
        path = OUT / book.filename
        book.finish(path)
        built.append(path)
        print(f"  已產生 {path.relative_to(ROOT)}")
    if "--no-verify" in sys.argv or not built:
        return
    sys.exit(0 if verify(built) else 1)


if __name__ == "__main__":
    main()
